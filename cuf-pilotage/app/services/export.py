"""
Service d'export Excel enrichi — Rapport mensuel Chaîne 4, Scierie CUF.
6 feuilles : Résumé | Production | TRS | Analyse Essence | Maintenance | Pareto
"""
import io
from collections import defaultdict
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Parametre
from app.services.trs import (
    pareto_arrets, calcule_pertes_equipe, _prix_production,
    manque_a_gagner_agrege, calcule_manque_gagner,
)

VERT_FONCE  = "1B5E20"
VERT_CLAIR  = "C8E6C9"
ORANGE_PALE = "FFE0B2"
ROUGE_PALE  = "FFCDD2"
JAUNE_PALE  = "FFF9C4"
BLANC       = "FFFFFF"
GRIS_CLAIR  = "F5F5F5"
GRIS_TOTAL  = "E0E0E0"
BLEU_PALE   = "BBDEFB"

NOMS_MOIS = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
             'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']


def _style_entete(cell, bg_hex=VERT_FONCE, fg_hex=BLANC, taille=11, gras=True):
    cell.font = Font(bold=gras, color=fg_hex, size=taille, name='Calibri')
    cell.fill = PatternFill("solid", fgColor=bg_hex)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='BBBBBB')
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _style_cellule(cell, gras=False, couleur_bg=None, align='center'):
    cell.font = Font(bold=gras, size=10, name='Calibri')
    cell.alignment = Alignment(horizontal=align, vertical='center')
    if couleur_bg:
        cell.fill = PatternFill("solid", fgColor=couleur_bg)
    thin = Side(style='thin', color='DDDDDD')
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _bg_trs(valeur):
    if valeur is None:
        return GRIS_CLAIR
    if valeur >= 70:
        return VERT_CLAIR
    if valeur >= 50:
        return JAUNE_PALE
    return ROUGE_PALE


def _ligne_total(ws, row, vals_par_col, label='TOTAL', col_label=1):
    medium = Side(style='medium', color='9E9E9E')
    border = Border(left=medium, right=medium, top=medium, bottom=medium)
    max_col = max(vals_par_col.keys()) if vals_par_col else col_label
    for col in range(col_label, max_col + 1):
        val = label if col == col_label else vals_par_col.get(col)
        c = ws.cell(row=row, column=col, value=val)
        c.font = Font(bold=True, size=10, name='Calibri')
        c.fill = PatternFill("solid", fgColor=GRIS_TOTAL)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border
    ws.row_dimensions[row].height = 20


def _prix_manquants(equipes):
    for e in equipes:
        for p in e.productions:
            if (p.volume_conforme + p.volume_declass) > 0 and _prix_production(p) == 0:
                return True
    return False


# ── Feuille 1 : Résumé ───────────────────────────────────────────────────────

def _feuille_resume(wb, equipes, mois, annee, nb_brouillons=0):
    ws = wb.active
    ws.title = "Résumé"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 16

    ws.merge_cells('A1:C1')
    t = ws['A1']
    t.value = "RAPPORT MENSUEL — CHAÎNE 4, CUF EBOLOWA"
    t.font = Font(bold=True, size=14, color=BLANC, name='Calibri')
    t.fill = PatternFill("solid", fgColor=VERT_FONCE)
    t.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    ws.merge_cells('A2:C2')
    st = ws['A2']
    st.value = f"{NOMS_MOIS[mois]} {annee}"
    st.font = Font(bold=True, size=12, color="333333", name='Calibri')
    st.alignment = Alignment(horizontal='center', vertical='center')
    st.fill = PatternFill("solid", fgColor=VERT_CLAIR)
    ws.row_dimensions[2].height = 22

    if not equipes:
        ws['A4'] = "Aucune équipe enregistrée pour cette période."
        return

    trs_vals = [e.trs_global for e in equipes if e.trs_global is not None]
    trs_moyen = round(sum(trs_vals) / len(trs_vals), 1) if trs_vals else 0
    production_reelle = round(sum(e.volume_sorti for e in equipes), 1)
    objectif_m3 = float(Parametre.get('objectif_m3', 12.5))
    production_cible = objectif_m3 * len(equipes)
    ecart = round(production_cible - production_reelle, 1)

    vol_entree_total = round(sum(e.volume_entree for e in equipes), 1)
    rendement_global = round(production_reelle / vol_entree_total * 100, 1) if vol_entree_total > 0 else 0
    duree_totale_arrets = sum(e.duree_totale_arrets for e in equipes)

    pertes_list = [calcule_pertes_equipe(e) for e in equipes]
    total_attribution = sum(p['total_attribution'] for p in pertes_list)
    total_d = sum(p['perte_d'] for p in pertes_list)
    total_p = sum(p['perte_p'] for p in pertes_list)
    total_q = sum(p['perte_q'] for p in pertes_list)

    # P11 — Manque à gagner estimé agrégé (indicateur principal)
    manque_periode = manque_a_gagner_agrege(equipes)

    def pct(val, tot):
        return f"{round(val / tot * 100, 1)}%" if tot > 0 else "—"

    # Section A — KPIs
    row = 4
    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'].value = "A — INDICATEURS CLÉS"
    _style_entete(ws[f'A{row}'], bg_hex="37474F")
    ws.row_dimensions[row].height = 22
    row += 1

    kpis = [
        ("Équipes analysées (validées chef + verrouillées)", len(equipes), None),
        ("TRS moyen", f"{trs_moyen}%", _bg_trs(trs_moyen)),
        ("Benchmark international (cible TRS)", "≥ 60%", None),
        ("Production réelle", f"{production_reelle} m³", None),
        ("Production objectif", f"{production_cible} m³", None),
        ("Écart production (non réalisé)", f"{ecart} m³", ROUGE_PALE if ecart > 0 else VERT_CLAIR),
        ("Rendement matière global", f"{rendement_global}%", _bg_trs(rendement_global)),
        ("Durée totale des arrêts", f"{duree_totale_arrets} min", None),
        ("Date de génération", date.today().strftime("%d/%m/%Y"), None),
    ]
    if nb_brouillons > 0:
        kpis.insert(1, (f"Fiches non validées ({nb_brouillons})",
                        "Non inclus dans ce rapport", ORANGE_PALE))

    for label, valeur, force_bg in kpis:
        bg = force_bg if force_bg else (GRIS_CLAIR if row % 2 == 0 else BLANC)
        c_l = ws.cell(row=row, column=1, value=label)
        ws.merge_cells(f'B{row}:C{row}')
        c_v = ws.cell(row=row, column=2, value=valeur)
        _style_cellule(c_l, gras=True, couleur_bg=bg, align='left')
        _style_cellule(c_v, couleur_bg=bg)
        ws.row_dimensions[row].height = 22
        row += 1

    row += 1

    # Section B — Manque à gagner estimé (P11 — indicateur principal)
    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'].value = "B — MANQUE À GAGNER ESTIMÉ (FCFA)"
    _style_entete(ws[f'A{row}'], bg_hex="B71C1C")
    ws.row_dimensions[row].height = 22
    row += 1

    for j, hdr in enumerate(['Composante', 'Montant (FCFA)', 'Note'], start=1):
        _style_entete(ws.cell(row=row, column=j, value=hdr), bg_hex="37474F")
    ws.row_dimensions[row].height = 20
    row += 1

    for label, val, note, bg in [
        ("Valeur potentielle (objectif × prix moyen)",
         int(manque_periode['valeur_potentielle']), "Cible théorique",        BLANC),
        ("Valeur réelle valorisée (conforme + déclassé)",
         int(manque_periode['valeur_reelle_valorisee']), "Production effective", VERT_CLAIR),
        ("Manque à gagner estimé",
         int(manque_periode['manque_a_gagner_estime']), "Potentielle − réelle", ROUGE_PALE),
    ]:
        _style_cellule(ws.cell(row=row, column=1, value=label), couleur_bg=bg, align='left')
        _style_cellule(ws.cell(row=row, column=2, value=val), couleur_bg=bg)
        _style_cellule(ws.cell(row=row, column=3, value=note), couleur_bg=bg)
        ws.row_dimensions[row].height = 20
        row += 1

    row += 1

    # Section B-bis — Causes probables D/P/Q (attribution causale, secondaire)
    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'].value = "B-bis — CAUSES PROBABLES (attribution D / P / Q)"
    _style_entete(ws[f'A{row}'], bg_hex="37474F")
    ws.row_dimensions[row].height = 22
    row += 1

    for j, hdr in enumerate(['Cause', 'Montant (FCFA)', '% du total'], start=1):
        _style_entete(ws.cell(row=row, column=j, value=hdr), bg_hex="546E7A")
    ws.row_dimensions[row].height = 20
    row += 1

    for label, val, bg in [
        ("Cause D — Arrêts non planifiés", total_d, ROUGE_PALE),
        ("Cause P — Sous-performance",     total_p, ORANGE_PALE),
        ("Cause Q — Qualité matière",      total_q, JAUNE_PALE),
    ]:
        _style_cellule(ws.cell(row=row, column=1, value=label), couleur_bg=bg, align='left')
        _style_cellule(ws.cell(row=row, column=2, value=int(val)), couleur_bg=bg)
        _style_cellule(ws.cell(row=row, column=3, value=pct(val, total_attribution)), couleur_bg=bg)
        ws.row_dimensions[row].height = 20
        row += 1

    # Note d'avertissement sur les chevauchements
    ws.merge_cells(f'A{row}:C{row}')
    note = ws[f'A{row}']
    note.value = ("Attribution causale indicative — certaines causes peuvent se "
                  "chevaucher. Indicateur principal : manque à gagner ci-dessus.")
    note.font = Font(italic=True, size=9, color="616161", name='Calibri')
    note.fill = PatternFill("solid", fgColor=GRIS_CLAIR)
    note.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 28
    row += 2

    # Section C — Pareto top 5
    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'].value = "C — TOP 5 CAUSES D'ARRÊT (Pareto)"
    _style_entete(ws[f'A{row}'], bg_hex="4A148C")
    ws.row_dimensions[row].height = 22
    row += 1

    for j, hdr in enumerate(['Cause', 'Durée (min)', '% cumulé'], start=1):
        _style_entete(ws.cell(row=row, column=j, value=hdr), bg_hex="37474F")
    ws.row_dimensions[row].height = 20
    row += 1

    pareto = pareto_arrets(equipes)
    for i, item in enumerate(pareto[:5]):
        bg = ROUGE_PALE if i == 0 else (ORANGE_PALE if i == 1 else BLANC)
        _style_cellule(ws.cell(row=row, column=1, value=item['cause']), couleur_bg=bg, align='left')
        _style_cellule(ws.cell(row=row, column=2, value=item['duree']), couleur_bg=bg)
        _style_cellule(ws.cell(row=row, column=3, value=f"{item['pct_cumule']}%"), couleur_bg=bg)
        ws.row_dimensions[row].height = 20
        row += 1

    # Section D — Alerte prix manquants
    if _prix_manquants(equipes):
        row += 1
        ws.merge_cells(f'A{row}:C{row}')
        alerte = ws[f'A{row}']
        alerte.value = ("ALERTE : Des prix de vente sont manquants — les pertes financières "
                        "sont sous-estimées. Saisir les prix dans Paramètres.")
        alerte.font = Font(bold=True, size=10, color=BLANC, name='Calibri')
        alerte.fill = PatternFill("solid", fgColor="B71C1C")
        alerte.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[row].height = 32


# ── Feuille 2 : Production ───────────────────────────────────────────────────

def _feuille_production(wb, equipes):
    ws = wb.create_sheet("Production")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = '1:1'

    entetes = [
        'Date', 'Équipe', 'Statut', 'Essences', 'Effectif',
        'Entrée (m³)', 'Conforme (m³)', 'Déclassé (m³)', 'Déchets (m³)',
        'Rendement %', 'Arrêts (min)',
        'TRS (%)', 'Dispo (%)', 'Perf (%)', 'Qualité (%)',
        'Perte D (FCFA)', 'Perte P (FCFA)', 'Perte Q décl', 'Perte Q déch',
        'Perte Q tot', 'Attribution D/P/Q (FCFA)',
        'Soumis le', 'Modifié par', 'Notes',
    ]
    largeurs = [14, 12, 14, 20, 10, 12, 12, 11, 11, 12, 12,
                10, 10, 10, 12, 15, 15, 14, 14, 13, 16, 14, 14, 30]

    for j, (h, w) in enumerate(zip(entetes, largeurs), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
        _style_entete(ws.cell(row=1, column=j, value=h))
    ws.row_dimensions[1].height = 30
    ws.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}1"

    BG_STATUT = {
        'verrouille': VERT_CLAIR,
        'valide_chef': VERT_CLAIR,
        'soumis': JAUNE_PALE,
        'a_verifier': BLEU_PALE,
        'a_corriger': ORANGE_PALE,
        'brouillon': ORANGE_PALE,
    }

    for i, e in enumerate(equipes, start=2):
        pertes = calcule_pertes_equipe(e)
        rendement = round(e.volume_sorti / e.volume_entree * 100, 1) if e.volume_entree > 0 else 0
        soumis_le = e.soumis_le.strftime('%d/%m/%Y') if e.soumis_le else ''
        valeurs = [
            e.date.strftime('%d/%m/%Y') if e.date else '',
            e.numero_equipe, e.statut, e.essences_label, e.effectif,
            e.volume_entree, e.volume_conforme, e.volume_declass, e.volume_dechets,
            rendement, e.duree_totale_arrets,
            e.trs_global, e.trs_disponibilite, e.trs_performance, e.trs_qualite,
            pertes['perte_d'], pertes['perte_p'],
            pertes['perte_q_declass'], pertes['perte_q_dechets'],
            pertes['perte_q'], pertes['total'],
            soumis_le, e.modifie_par or '', e.notes or '',
        ]
        bg_ligne = BLANC if i % 2 == 0 else GRIS_CLAIR
        for j, val in enumerate(valeurs, start=1):
            if j == 3:
                bg = BG_STATUT.get(val, bg_ligne)
            elif j in (10, 12, 13, 14, 15):
                bg = _bg_trs(val)
            elif j in (16, 17, 18, 19, 20, 21):
                bg = ROUGE_PALE if (val or 0) > 0 else bg_ligne
            else:
                bg = bg_ligne
            _style_cellule(ws.cell(row=i, column=j, value=val),
                           couleur_bg=bg, align='left' if j in (4, 24) else 'center')
        ws.row_dimensions[i].height = 18

    if equipes:
        pertes_all = [calcule_pertes_equipe(e) for e in equipes]
        tot_row = len(equipes) + 2
        _ligne_total(ws, tot_row, {
            6:  round(sum(e.volume_entree for e in equipes), 1),
            7:  round(sum(e.volume_conforme for e in equipes), 1),
            8:  round(sum(e.volume_declass for e in equipes), 1),
            9:  round(sum(e.volume_dechets for e in equipes), 1),
            11: sum(e.duree_totale_arrets for e in equipes),
            16: int(sum(p['perte_d'] for p in pertes_all)),
            17: int(sum(p['perte_p'] for p in pertes_all)),
            18: int(sum(p['perte_q_declass'] for p in pertes_all)),
            19: int(sum(p['perte_q_dechets'] for p in pertes_all)),
            20: int(sum(p['perte_q'] for p in pertes_all)),
            21: int(sum(p['total'] for p in pertes_all)),
        })


# ── Feuille 3 : TRS ──────────────────────────────────────────────────────────

def _feuille_trs(wb, equipes):
    ws = wb.create_sheet("TRS")
    ws.sheet_view.showGridLines = False

    # Tableau A — synthèse créneaux
    entetes_a = ['Créneau', 'Nb postes', 'TRS moy (%)', 'Dispo moy (%)',
                 'Perf moy (%)', 'Qualité moy (%)', 'Prod totale (m³)',
                 'Arrêts tot (min)', 'Manque à gagner (FCFA)']
    largeurs_a = [18, 12, 14, 14, 14, 14, 16, 16, 18]

    ws.merge_cells('A1:I1')
    ws['A1'].value = "SYNTHÈSE PAR CRÉNEAU"
    _style_entete(ws['A1'], bg_hex="37474F")
    ws.row_dimensions[1].height = 22

    for j, (h, w) in enumerate(zip(entetes_a, largeurs_a), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
        _style_entete(ws.cell(row=2, column=j, value=h))
    ws.row_dimensions[2].height = 22

    groupes = [
        ('Matin',      [e for e in equipes if e.numero_equipe == 'Matin']),
        ('Après-midi', [e for e in equipes if e.numero_equipe == 'Apres-midi']),
    ]

    def moy(lst):
        return round(sum(lst) / len(lst), 1) if lst else 0

    row_a = 3
    for creneau, groupe in groupes:
        if not groupe:
            continue
        trs_v   = [e.trs_global        for e in groupe if e.trs_global is not None]
        dispo_v = [e.trs_disponibilite for e in groupe if e.trs_disponibilite is not None]
        perf_v  = [e.trs_performance   for e in groupe if e.trs_performance is not None]
        qual_v  = [e.trs_qualite       for e in groupe if e.trs_qualite is not None]
        prod_t  = round(sum(e.volume_sorti for e in groupe), 1)
        arr_t   = sum(e.duree_totale_arrets for e in groupe)
        per_t   = int(sum(calcule_manque_gagner(e)['manque_a_gagner_estime'] for e in groupe))
        trs_m   = moy(trs_v)
        vals = [creneau, len(groupe), trs_m, moy(dispo_v),
                moy(perf_v), moy(qual_v), prod_t, arr_t, per_t]
        bg = _bg_trs(trs_m)
        for j, val in enumerate(vals, start=1):
            _style_cellule(ws.cell(row=row_a, column=j, value=val),
                           couleur_bg=bg, align='left' if j == 1 else 'center')
        ws.row_dimensions[row_a].height = 20
        row_a += 1

    # Tableau B — détail journalier
    row_b_title = row_a + 2
    ws.merge_cells(f'A{row_b_title}:I{row_b_title}')
    ws[f'A{row_b_title}'].value = "DÉTAIL JOURNALIER"
    _style_entete(ws[f'A{row_b_title}'], bg_hex="37474F")
    ws.row_dimensions[row_b_title].height = 22

    entetes_b = ['Date', 'Équipe', 'TRS (%)', 'Dispo (%)', 'Perf (%)',
                 'Qualité (%)', 'Prod (m³)', 'Arrêts (min)', 'Manque à gagner (FCFA)']
    row_b_hdr = row_b_title + 1
    for j, h in enumerate(entetes_b, start=1):
        _style_entete(ws.cell(row=row_b_hdr, column=j, value=h))
    ws.row_dimensions[row_b_hdr].height = 22

    ws.freeze_panes = f'A{row_b_hdr + 1}'
    ws.auto_filter.ref = f"A{row_b_hdr}:I{row_b_hdr}"

    for i, e in enumerate(sorted(equipes, key=lambda x: x.date)):
        row = row_b_hdr + 1 + i
        per_tot = int(calcule_manque_gagner(e)['manque_a_gagner_estime'])
        vals = [
            e.date.strftime('%d/%m/%Y') if e.date else '',
            e.numero_equipe, e.trs_global, e.trs_disponibilite,
            e.trs_performance, e.trs_qualite,
            round(e.volume_sorti, 1), e.duree_totale_arrets, per_tot,
        ]
        bg_ligne = BLANC if i % 2 == 0 else GRIS_CLAIR
        for j, val in enumerate(vals, start=1):
            bg = _bg_trs(val) if j in (3, 4, 5, 6) else bg_ligne
            _style_cellule(ws.cell(row=row, column=j, value=val),
                           couleur_bg=bg, align='left' if j == 1 else 'center')
        ws.row_dimensions[row].height = 18


# ── Feuille 4 : Analyse Essence ──────────────────────────────────────────────

def _feuille_essence(wb, equipes):
    ws = wb.create_sheet("Essence")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'

    entetes = [
        'Essence', 'Nb postes', 'Vol entré (m³)', 'Vol sorti (m³)',
        'Rendement %', 'Prix moy (FCFA/m³)',
        'Vol déclassé (m³)', 'Vol déchets (m³)',
        'Perte Q décl (FCFA)', 'Perte Q déch (FCFA)', 'Perte Q totale (FCFA)',
    ]
    largeurs = [18, 11, 16, 16, 14, 18, 16, 16, 18, 18, 18]

    for j, (h, w) in enumerate(zip(entetes, largeurs), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
        _style_entete(ws.cell(row=1, column=j, value=h))
    ws.row_dimensions[1].height = 30
    ws.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}1"

    BG_ESSENCE = {
        'Ayous':    JAUNE_PALE,
        'Azobé':    'FFCCBC',
        'Iroko':    VERT_CLAIR,
        'Movingui': BLEU_PALE,
    }

    taux_revente = float(Parametre.get('taux_revente_rebut', 0.70))
    data = defaultdict(lambda: {
        'vol_entree': 0.0, 'vol_sorti': 0.0,
        'vol_declass': 0.0, 'vol_dechets': 0.0,
        'vol_prix': 0.0,
        'perte_q_declass': 0.0, 'perte_q_dechets': 0.0,
    })
    postes_ess = defaultdict(set)

    for e in equipes:
        for p in e.productions:
            ess = p.essence
            prix = _prix_production(p)
            vol_p = p.volume_conforme + p.volume_declass
            postes_ess[ess].add(e.id)
            data[ess]['vol_entree']      += p.volume_entree
            data[ess]['vol_sorti']       += vol_p
            data[ess]['vol_declass']     += p.volume_declass
            data[ess]['vol_dechets']     += p.volume_dechets
            data[ess]['vol_prix']        += prix * vol_p
            data[ess]['perte_q_declass'] += p.volume_declass * prix * (1 - taux_revente)
            data[ess]['perte_q_dechets'] += p.volume_dechets * prix

    row = 2
    for ess in sorted(data.keys()):
        d = data[ess]
        ve, vs = d['vol_entree'], d['vol_sorti']
        rendement = round(vs / ve * 100, 1) if ve > 0 else 0
        prix_moy = int(round(d['vol_prix'] / vs, 0)) if vs > 0 else 0
        pq_dec = int(d['perte_q_declass'])
        pq_dch = int(d['perte_q_dechets'])
        pq_tot = pq_dec + pq_dch

        bg_ess = BG_ESSENCE.get(ess, BLANC)
        if rendement >= 60:
            bg_rend = VERT_CLAIR
        elif rendement >= 40:
            bg_rend = JAUNE_PALE
        else:
            bg_rend = ROUGE_PALE

        valeurs = [
            ess, len(postes_ess[ess]),
            round(ve, 1), round(vs, 1),
            rendement, prix_moy,
            round(d['vol_declass'], 1), round(d['vol_dechets'], 1),
            pq_dec, pq_dch, pq_tot,
        ]
        for j, val in enumerate(valeurs, start=1):
            if j == 5:
                bg = bg_rend
            elif j == 6:
                bg = ROUGE_PALE if prix_moy == 0 else bg_ess
            elif j in (9, 10, 11):
                bg = ROUGE_PALE if val > 0 else bg_ess
            else:
                bg = bg_ess
            _style_cellule(ws.cell(row=row, column=j, value=val),
                           couleur_bg=bg, align='left' if j == 1 else 'center')
        ws.row_dimensions[row].height = 20
        row += 1


# ── Feuille 5 : Maintenance ──────────────────────────────────────────────────

BG_CAT = {
    'Mécanique':              ROUGE_PALE,
    'Organisationnelle':      ORANGE_PALE,
    'Maintenance planifiée':  JAUNE_PALE,
    'Électrique':             BLEU_PALE,
}


def _feuille_maintenance(wb, equipes):
    ws = wb.create_sheet("Maintenance")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = '1:1'

    entetes = [
        'Date', 'Équipe', 'Machine', 'Heure début', 'Heure fin',
        'Durée (min)', 'Cause', 'Catégorie', 'Planifié ?', 'Impact TRS', 'Notes',
    ]
    largeurs = [14, 12, 18, 12, 10, 12, 40, 22, 12, 12, 30]

    for j, (h, w) in enumerate(zip(entetes, largeurs), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
        _style_entete(ws.cell(row=1, column=j, value=h), bg_hex="37474F")
    ws.row_dimensions[1].height = 28
    ws.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}1"

    ligne = 2
    total_incidents = 0
    total_duree = 0

    for e in sorted(equipes, key=lambda x: x.date):
        for a in e.arrets:
            planifie = 'Oui' if a.categorie == 'Maintenance planifiée' else 'Non'
            impact = 'Non' if a.categorie == 'Maintenance planifiée' else 'Oui'
            valeurs = [
                e.date.strftime('%d/%m/%Y') if e.date else '',
                e.numero_equipe, a.machine,
                a.heure_debut, a.heure_fin, a.duree_min,
                a.cause, a.categorie,
                planifie, impact, a.notes or '',
            ]
            bg_row = BG_CAT.get(a.categorie, BLANC)
            for j, val in enumerate(valeurs, start=1):
                if j == 9:
                    cell_bg = VERT_CLAIR if val == 'Oui' else ROUGE_PALE
                elif j == 10:
                    cell_bg = ORANGE_PALE if val == 'Oui' else BLANC
                else:
                    cell_bg = bg_row
                _style_cellule(ws.cell(row=ligne, column=j, value=val),
                               couleur_bg=cell_bg,
                               align='left' if j in (7, 8, 11) else 'center')
            ws.row_dimensions[ligne].height = 18
            total_incidents += 1
            total_duree += a.duree_min or 0
            ligne += 1

    if ligne == 2:
        ws.cell(row=2, column=1, value="Aucun arrêt enregistré pour cette période.")
    else:
        _ligne_total(ws, ligne, {6: total_duree},
                     label=f'TOTAL ({total_incidents} incidents)')


# ── Feuille 6 : Pareto ───────────────────────────────────────────────────────

def _feuille_pareto(wb, equipes):
    ws = wb.create_sheet("Pareto")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'

    entetes = ['Cause', 'Durée (min)', 'Occurrences', 'Machines',
               'Catégories', '% total', '% cumulé']
    largeurs = [45, 14, 14, 30, 28, 12, 12]

    for j, (h, w) in enumerate(zip(entetes, largeurs), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
        _style_entete(ws.cell(row=1, column=j, value=h), bg_hex="4A148C")
    ws.row_dimensions[1].height = 28
    ws.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}1"

    cumul = defaultdict(lambda: {'duree': 0, 'count': 0,
                                  'machines': set(), 'categories': set()})
    for e in equipes:
        for a in e.arrets:
            if a.duree_min:
                cumul[a.cause]['duree']      += a.duree_min
                cumul[a.cause]['count']      += 1
                cumul[a.cause]['machines'].add(a.machine)
                cumul[a.cause]['categories'].add(a.categorie)

    total_duree = sum(v['duree'] for v in cumul.values())
    if total_duree == 0:
        ws.cell(row=2, column=1, value="Aucun arrêt enregistré pour cette période.")
        return

    resultats = sorted(
        [{'cause': k, **v} for k, v in cumul.items()],
        key=lambda x: x['duree'], reverse=True
    )

    cumul_pct = 0.0
    for r in resultats:
        r['pct'] = round(r['duree'] / total_duree * 100, 1)
        cumul_pct += r['pct']
        r['pct_cumule'] = round(cumul_pct, 1)

    row = 2
    for r in resultats:
        est_80 = r['pct_cumule'] <= 80.0
        premiere_cat = next(iter(r['categories']), '')
        bg = BG_CAT.get(premiere_cat, BLANC)
        valeurs = [
            r['cause'], r['duree'], r['count'],
            ', '.join(sorted(r['machines'])),
            ', '.join(sorted(r['categories'])),
            f"{r['pct']}%", f"{r['pct_cumule']}%",
        ]
        for j, val in enumerate(valeurs, start=1):
            c = ws.cell(row=row, column=j, value=val)
            c.font = Font(bold=est_80, size=10, name='Calibri')
            c.fill = PatternFill("solid", fgColor=bg)
            c.alignment = Alignment(
                horizontal='left' if j in (1, 4, 5) else 'center',
                vertical='center')
            thin = Side(style='thin', color='DDDDDD')
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.row_dimensions[row].height = 18
        row += 1


# ── Fonction principale ──────────────────────────────────────────────────────

def generer_rapport_excel(equipes, mois, annee, nb_brouillons=0):
    wb = Workbook()
    _feuille_resume(wb, equipes, mois, annee, nb_brouillons)
    _feuille_production(wb, equipes)
    _feuille_trs(wb, equipes)
    _feuille_essence(wb, equipes)
    _feuille_maintenance(wb, equipes)
    _feuille_pareto(wb, equipes)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
