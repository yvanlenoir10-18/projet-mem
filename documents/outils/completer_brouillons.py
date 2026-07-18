#!/usr/bin/env python3
"""
Crée des fiches BROUILLON pour les quarts manquants de la période réelle.
- Ne remplit AUCUNE valeur de production/arrêt inventée.
- Statut 'brouillon' => exclu de tous les KPI (TRS, Pareto...), donc ne fausse rien.
- En note : indice des essences vues ce jour dans la base de production (feuille 8),
  là où l'info existe, pour aider l'opérateur à compléter.
"""
import os
from datetime import timedelta
from collections import defaultdict
import openpyxl

_RACINE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # racine du dépôt (documents/outils/..)
import sys
sys.path.insert(0, os.path.join(_RACINE, 'cuf-pilotage'))
from app import create_app, db
from app.models import Equipe

XLSX = os.path.join(_RACINE, 'documents', 'collecte', 'Suivi_bicoupe_CUF_mai-juin-2026.xlsx')
CRENEAUX = ['Matin', 'Apres-midi', 'Nuit']
USER_ID = 1  # saisie@cuf.cm


def essences_par_date():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb['8-PRODUCTION POSTE']
    m = defaultdict(set)
    for r in range(5, ws.max_row + 1):
        d = ws.cell(r, 1).value
        e = ws.cell(r, 4).value
        if hasattr(d, 'date') and e:
            m[d.date()].add(str(e).strip())
    return m


def main():
    ess_map = essences_par_date()
    app = create_app()
    with app.app_context():
        # purge d'éventuels brouillons auto déjà créés AVANT de calculer les
        # quarts existants (sinon, en ré-exécution, les anciens brouillons
        # masqueraient les trous et rien ne serait recréé).
        anciens = Equipe.query.filter_by(statut='brouillon',
                                         operateur_nom='À compléter').all()
        for e in anciens:
            db.session.delete(e)
        db.session.commit()

        qs = Equipe.query.all()
        existants = {(e.date, e.numero_equipe) for e in qs}
        dmin = min(e.date for e in qs)
        dmax = max(e.date for e in qs)

        cree = 0
        d = dmin
        while d <= dmax:
            for c in CRENEAUX:
                if (d, c) in existants:
                    continue
                hint = ess_map.get(d)
                note = ("Brouillon généré automatiquement — quart sans relevé TRS "
                        "dans la base terrain. À compléter par l'opérateur.")
                if hint:
                    note += (" Essences produites ce jour (tous postes confondus, "
                             "source base production, à répartir) : "
                             + ", ".join(sorted(hint)) + ".")
                eq = Equipe(
                    date=d, numero_equipe=c, effectif=10,
                    statut='brouillon',
                    operateur_nom='À compléter',
                    rempli_par_nom='À compléter',
                    mode_saisie='directe',
                    notes=note,
                    user_id=USER_ID,
                )
                db.session.add(eq)
                cree += 1
            d += timedelta(days=1)
        db.session.commit()

        total = Equipe.query.count()
        verrou = Equipe.query.filter_by(statut='verrouille').count()
        brouill = Equipe.query.filter_by(statut='brouillon').count()
        print(f'Brouillons créés : {cree}')
        print(f'Base : {total} fiches -> {verrou} remplies (verrouillées) · {brouill} brouillons à compléter')
        print(f'Période couverte en continu : {dmin} -> {dmax}')


if __name__ == '__main__':
    main()
